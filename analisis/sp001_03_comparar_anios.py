"""Actividad 3: comparar la estructura de columnas entre 2023, 2024 y 2025.

Uso:  python sp001_03_comparar_anios.py
"""
import pandas as pd
from openpyxl import load_workbook

from utils import ANIOS, COLUMNAS_USADAS, FUENTES, cargar, guardar, llave, ruta_excel


def encabezados(fuente, anio):
    """Lee solo la primera fila del Excel (rápido)."""
    wb = load_workbook(ruta_excel(fuente, anio), read_only=True)
    fila = next(wb[wb.sheetnames[0]].iter_rows(max_row=1, values_only=True))
    wb.close()
    return list(fila)


# 1. Nombres y orden de columnas
print("### NOMBRES Y ORDEN DE COLUMNAS")
for fuente in FUENTES:
    h = {anio: encabezados(fuente, anio) for anio in ANIOS}
    iguales = h[2023] == h[2024] == h[2025]
    print(f"  {fuente:15s} columnas por año: {[len(h[a]) for a in ANIOS]}  identicas: {iguales}")
    if not iguales:
        for anio in ANIOS:
            print(f"    {anio} difiere en: {set(h[anio]) ^ set(h[2025])}")

# 2. Tipos de las columnas usadas
tipos = {}
for anio in ANIOS:
    for fuente, columnas in COLUMNAS_USADAS.items():
        df = cargar(fuente, anio)
        for col in columnas:
            tipos.setdefault((fuente, col), {})[anio] = str(df[col].dtype)
tabla_tipos = pd.DataFrame(
    [{"fuente": f, "columna": c, **t} for (f, c), t in tipos.items()]
)
tabla_tipos["igual_en_3_anios"] = tabla_tipos[ANIOS].nunique(axis=1) == 1
print("\n### TIPOS DE LAS COLUMNAS USADAS")
print(tabla_tipos.to_string(index=False))

# 3. Valores y calidad por año
print("\n### VALORES Y CALIDAD POR AÑO")
for anio in ANIOS:
    adj = cargar("ADJUDICACIONES", anio)
    con = cargar("CONVOCATORIAS", anio)
    pos = cargar("POSTOR", anio)
    cod = con["codigoitem"].dropna()
    repetidos = pos.assign(_k=llave(pos, "POSTOR")).duplicated(["_k", "ruc_codigo_postor"]).sum()
    print(f"\n  {anio}")
    print(f"    moneda:            {sorted(adj['moneda'].unique())}")
    print(f"    objetocontractual: {sorted(adj['objetocontractual'].unique())}")
    print(f"    codigoitem lleno:  {100 * con['codigoitem'].notna().mean():.1f}%"
          f"   con decimales distintos de .0: {int((cod % 1 != 0).sum())}")
    print(f"    postores: filas con el mismo RUC repetido en el mismo item: {int(repetidos)}")

guardar(tabla_tipos, "sp001/03_tipos_por_anio.csv")
