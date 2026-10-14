import csv
import math
import random
import statistics
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from agents import market_analyst  # noqa: E402

TESTS_DIR = Path(__file__).parent
DATA_DIR = TESTS_DIR.parent / "data"
SEED = 2026
TOLERANCE = 0.000001
STRATA = [("5 a 9", 5, 9), ("10 a 29", 10, 29), ("30 o más", 30, math.inf)]
GROUPS_PER_STRATUM = 10
LA_LIBERTAD = "LA LIBERTAD"


def read_csv(name):
    with open(DATA_DIR / name, encoding="utf-8", newline="") as file:
        return list(csv.DictReader(file))


def linear_percentile(sorted_values, p):
    """Interpolación lineal entre posiciones (definición de PERCENTIL.INC de Excel)."""
    position = (len(sorted_values) - 1) * p
    lower = math.floor(position)
    upper = min(lower + 1, len(sorted_values) - 1)
    fraction = position - lower
    return sorted_values[lower] + fraction * (sorted_values[upper] - sorted_values[lower])


def control_evidence(n):
    """Mínimos de SP-002 escritos de nuevo, sin importar rules.py."""
    if n < 5:
        return "no_evidence"
    if n < 10:
        return "insufficient_evidence"
    return "with_verdict"


def equal(a, b):
    """Iguales dentro de la tolerancia; dos valores ausentes también son iguales."""
    if a is None or b is None:
        return a is None and b is None
    return abs(a - b) <= TOLERANCE


catalog = read_csv("catalog.csv")
awards = read_csv("awards.csv")

rng = random.Random(SEED)
selection = []
for stratum_name, low, high in STRATA:
    in_stratum = [g for g in catalog if low <= int(g["n_adjudicaciones"]) <= high]
    for good in sorted(rng.sample(in_stratum, GROUPS_PER_STRATUM), key=lambda g: int(g["n_adjudicaciones"])):
        selection.append((stratum_name, good))

rows, data_by_group = [], []
for number, (stratum, good) in enumerate(selection, start=1):
    code, unit = int(good["codigoitem"]), good["unidad_medida"]

    own = [a for a in awards if int(a["codigoitem"]) == code and a["unidad_medida"] == unit]
    prices = sorted(float(a["precio_unitario"]) for a in own)
    years = [int(a["anio"]) for a in own]
    bidders = sorted(int(a["n_postores"]) for a in own if a["n_postores"] != "")
    departments = [a["entidad_departamento"] for a in own]
    n_single = sum(1 for b in bidders if b == 1)

    control = {
        "n": len(prices),
        "period": f"{min(years)}-{max(years)}",
        "p25": linear_percentile(prices, 0.25),
        "median": linear_percentile(prices, 0.50),
        "p75": linear_percentile(prices, 0.75),
        "evidence": control_evidence(len(prices)),
        "n_with_bidders": len(bidders),
        "median_bidders": linear_percentile(bidders, 0.50) if bidders else None,
        "n_single_bidder": n_single,
        "pct_single_bidder": 100 * n_single / len(bidders) if bidders else None,
        "n_la_libertad": sum(1 for d in departments if d == LA_LIBERTAD),
    }
    q1, q2, q3 = statistics.quantiles(prices, n=4, method="inclusive")
    standard_median_bidders = statistics.median(bidders) if bidders else None
    data_by_group.append((code, prices, bidders, departments))

    app = market_analyst.analyze(code)

    row = {
        "group": number, "stratum": stratum, "code": code, "name": good["itemcubso"].strip(),
        "unit": unit,
        "n_app": app["n_awards"], "n_control": control["n"],
        "period_app": f"{app['year_from']}-{app['year_to']}", "period_control": control["period"],
        "evidence_app": app["evidence"], "evidence_control": control["evidence"],
    }

    differences = []
    for measure, standard in (("p25", q1), ("median", q2), ("p75", q3)):
        diff = abs(app[measure] - control[measure])
        differences += [diff, abs(app[measure] - standard)]
        row.update({
            f"{measure}_app": app[measure],
            f"{measure}_control": control[measure],
            f"{measure}_statistics": standard,
            f"{measure}_difference": diff,
        })
    row["max_price_difference"] = max(differences)
    prices_match = (
        row["n_app"] == row["n_control"]
        and row["period_app"] == row["period_control"]
        and row["evidence_app"] == row["evidence_control"]
        and row["max_price_difference"] <= TOLERANCE
    )

    for measure in ("n_with_bidders", "median_bidders", "n_single_bidder", "pct_single_bidder", "n_la_libertad"):
        row[f"{measure}_app"] = app[measure]
        row[f"{measure}_control"] = control[measure]
    row["median_bidders_statistics"] = standard_median_bidders
    competition_matches = (
        app["n_with_bidders"] == control["n_with_bidders"]
        and app["n_single_bidder"] == control["n_single_bidder"]
        and app["n_la_libertad"] == control["n_la_libertad"]
        and equal(app["median_bidders"], control["median_bidders"])
        and equal(app["median_bidders"], standard_median_bidders)
        and equal(app["pct_single_bidder"], control["pct_single_bidder"])
    )

    row["prices_match"] = prices_match
    row["competition_matches"] = competition_matches
    row["matches"] = prices_match and competition_matches
    rows.append(row)

with open(TESTS_DIR / "control_hu002.csv", "w", encoding="utf-8-sig", newline="") as file:
    writer = csv.DictWriter(file, fieldnames=list(rows[0]))
    writer.writeheader()
    writer.writerows(rows)

workbook = Workbook()
control_sheet = workbook.active
control_sheet.title = "Control"
prices_sheet = workbook.create_sheet("Precios")
bidders_sheet = workbook.create_sheet("Postores")
departments_sheet = workbook.create_sheet("Departamentos")
bold, blue = Font(name="Arial", bold=True), Font(name="Arial", color="0000FF")

for column, (code, prices, bidders, departments) in enumerate(data_by_group, start=1):
    for sheet, values in ((prices_sheet, prices), (bidders_sheet, bidders), (departments_sheet, departments)):
        sheet.cell(row=1, column=column, value=code).font = bold
        for excel_row, value in enumerate(values, start=2):
            sheet.cell(row=excel_row, column=column, value=value)

headers = [
    "Grupo", "Código", "Unidad",
    "N (Excel)", "P25 (Excel)", "Mediana (Excel)", "P75 (Excel)",
    "N (app)", "P25 (app)", "Mediana (app)", "P75 (app)", "Diferencia precios",
    "Con postores (Excel)", "Mediana postores (Excel)", "% postor único (Excel)", "La Libertad (Excel)",
    "Con postores (app)", "Mediana postores (app)", "% postor único (app)", "La Libertad (app)",
    "Diferencia competencia",
    "Coincide",
]
for column, text in enumerate(headers, start=1):
    control_sheet.cell(row=1, column=column, value=text).font = bold

last_data_row = max(len(p) for _, p, _, _ in data_by_group) + 1
for i, row in enumerate(rows, start=2):
    letter = get_column_letter(i - 1)
    prices_range = f"Precios!${letter}$2:${letter}${last_data_row}"
    bidders_range = f"Postores!${letter}$2:${letter}${last_data_row}"
    departments_range = f"Departamentos!${letter}$2:${letter}${last_data_row}"
    values = [
        row["group"], row["code"], row["unit"],
        f"=COUNT({prices_range})",
        f"=_xlfn.PERCENTILE.INC({prices_range},0.25)",
        f"=MEDIAN({prices_range})",
        f"=_xlfn.PERCENTILE.INC({prices_range},0.75)",
        row["n_app"], row["p25_app"], row["median_app"], row["p75_app"],
        f"=MAX(ABS(E{i}-I{i}),ABS(F{i}-J{i}),ABS(G{i}-K{i}))",
        f"=COUNT({bidders_range})",
        f'=IF(M{i}=0,"",MEDIAN({bidders_range}))',
        f'=IF(M{i}=0,"",COUNTIF({bidders_range},1)/M{i}*100)',
        f'=COUNTIF({departments_range},"{LA_LIBERTAD}")',
        row["n_with_bidders_app"],
        "" if row["median_bidders_app"] is None else row["median_bidders_app"],
        "" if row["pct_single_bidder_app"] is None else row["pct_single_bidder_app"],
        row["n_la_libertad_app"],
        f"=IF(M{i}=0,0,MAX(ABS(N{i}-R{i}),ABS(O{i}-S{i})))",
        f'=IF(AND(D{i}=H{i},L{i}<={TOLERANCE},M{i}=Q{i},P{i}=T{i},U{i}<={TOLERANCE}),"Sí","No")',
    ]
    for column, value in enumerate(values, start=1):
        cell = control_sheet.cell(row=i, column=column, value=value)
        if 8 <= column <= 11 or 17 <= column <= 20:
            cell.font = blue   
        if column in (5, 6, 7, 9, 10, 11):
            cell.number_format = "#,##0.000000"
        if column in (15, 19):
            cell.number_format = "0.000000"
        if column in (12, 21):
            cell.number_format = "0.0E+00"

last = len(rows) + 1
control_sheet.cell(row=last + 2, column=1, value="Discrepancias").font = bold
control_sheet.cell(row=last + 2, column=2, value=f'=COUNTIF(V2:V{last},"No")').font = bold
control_sheet.cell(row=last + 3, column=1,
                   value="Columnas en azul: valores calculados por el agente analista. "
                         "Columnas (Excel): fórmulas sobre las hojas Precios, Postores y Departamentos.")
widths = [8, 10, 11, 10, 16, 16, 16, 10, 16, 16, 16, 18, 14, 16, 16, 12, 14, 16, 16, 12, 18, 10]
for column, width in enumerate(widths, start=1):
    control_sheet.column_dimensions[get_column_letter(column)].width = width
for sheet in workbook.worksheets:
    for sheet_row in sheet.iter_rows():
        for cell in sheet_row:
            if cell.font.name != "Arial":
                cell.font = Font(name="Arial", bold=cell.font.bold, color=cell.font.color)
workbook.save(TESTS_DIR / "control_hu002.xlsx")

def show(value, fmt):
    return "—" if value is None else format(value, fmt)


print("### CONTROL HU-002 y HU-003: agente analista vs. herramientas independientes\n")
print(f"{'#':>2} {'estrato':<9} {'código':>7} {'N':>5} {'P25':>14} {'mediana':>14} {'P75':>14}"
      f"  {'post.':>5} {'único':>7} {'LL':>4}  resultado")
for r in rows:
    mark = "✔" if r["matches"] else "✘"
    print(f"{r['group']:>2} {r['stratum']:<9} {r['code']:>7} {r['n_app']:>5} "
          f"{r['p25_app']:>14,.4f} {r['median_app']:>14,.4f} {r['p75_app']:>14,.4f}  "
          f"{show(r['median_bidders_app'], '>5.1f')} {show(r['pct_single_bidder_app'], '>6.1f')}% "
          f"{r['n_la_libertad_app']:>4}  {mark}")

print(f"\nGrupos revisados: {len(rows)}")
print(f"Rango de precios (HU-002) con discrepancias: {sum(not r['prices_match'] for r in rows)}")
print(f"Competencia (HU-003) con discrepancias:      {sum(not r['competition_matches'] for r in rows)}")
print(f"Grupos con alguna discrepancia: {sum(not r['matches'] for r in rows)}  (criterio de aceptación: 0)")
print(f"Diferencia máxima en precios: {max(r['max_price_difference'] for r in rows):.2e} soles")
print(f"Adjudicaciones sin dato de postores en estos grupos: "
      f"{sum(r['n_app'] - r['n_with_bidders_app'] for r in rows)}")
print(f"\nDetalle: {TESTS_DIR / 'control_hu002.csv'}")
print(f"Libro para revisar en Excel: {TESTS_DIR / 'control_hu002.xlsx'}")